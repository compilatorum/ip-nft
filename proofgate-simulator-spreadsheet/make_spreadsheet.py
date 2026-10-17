"""make_spreadsheet.py
Gera um arquivo Excel proofgate_simulations.xlsx a partir dos CSVs na pasta proofgate-simulator-spreadsheet.
Dependências: pandas, openpyxl
Uso: python make_spreadsheet.py
"""

import os
import pandas as pd
from openpyxl import load_workbook

BASE_DIR = os.path.dirname(__file__)
OUT_XLSX = os.path.join(BASE_DIR, "proofgate_simulations.xlsx")

csv_files = [
    ("01_Index.csv", "01_Index"),
    ("02_Config_Mandate.csv", "02_Config_Mandate"),
    ("03_Market_Evidence.csv", "03_Market_Evidence"),
    ("04_Agent_Intent.csv", "04_Agent_Intent"),
    ("05_Scenarios.csv", "05_Scenarios"),
    ("06_Runs.csv", "06_Runs"),
    ("07_Invariants_Eval.csv", "07_Invariants_Eval"),
    ("08_Results_Summary.csv", "08_Results_Summary"),
    ("09_Health_Report.csv", "09_Health_Report"),
    ("10_Solana_Logs.csv", "10_Solana_Logs"),
]

with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
    for file_name, sheet_name in csv_files:
        path = os.path.join(BASE_DIR, file_name)
        if os.path.exists(path):
            df = pd.read_csv(path, dtype=str)
            df.to_excel(writer, sheet_name=sheet_name, index=False)

wb = load_workbook(OUT_XLSX)
ws_runs = wb["06_Runs"]
max_row = ws_runs.max_row

# Fórmulas iniciais básicas na aba Runs
for r in range(2, max_row + 1):
    ws_runs[f"I{r}"].value = f"=IFERROR(INDEX('03_Market_Evidence'!$L:$L, MATCH($F{r}, '03_Market_Evidence'!$A:$A, 0)), \"\")"
    ws_runs[f"J{r}"].value = f"=IFERROR(INDEX('03_Market_Evidence'!$M:$M, MATCH($F{r}, '03_Market_Evidence'!$A:$A, 0)), \"\")"
    ws_runs[f"K{r}"].value = f"=IFERROR(INDEX('03_Market_Evidence'!$N:$N, MATCH($F{r}, '03_Market_Evidence'!$A:$A, 0)), \"\")"
    ws_runs[f"L{r}"].value = f"=IFERROR(INDEX('03_Market_Evidence'!$G:$G, MATCH($F{r}, '03_Market_Evidence'!$A:$A, 0)), \"\")"
    ws_runs[f"M{r}"].value = f"=IFERROR(IF(SEARCH(INDEX('04_Agent_Intent'!$H:$H, MATCH($G{r}, '04_Agent_Intent'!$A:$A, 0)), INDEX('02_Config_Mandate'!$B:$B, MATCH($E{r}, '02_Config_Mandate'!$A:$A, 0)))>0, TRUE, FALSE), FALSE)"
    ws_runs[f"N{r}"].value = f"=IF(AND(INDEX('04_Agent_Intent'!$C:$C, MATCH($G{r}, '04_Agent_Intent'!$A:$A, 0)) <= INDEX('02_Config_Mandate'!$C:$C, MATCH($E{r}, '02_Config_Mandate'!$A:$A, 0)), (INDEX('04_Agent_Intent'!$C:$C, MATCH($G{r}, '04_Agent_Intent'!$A:$A, 0)) + INDEX('04_Agent_Intent'!$E:$E, MATCH($G{r}, '04_Agent_Intent'!$A:$A, 0))) <= INDEX('02_Config_Mandate'!$D:$D, MATCH($E{r}, '02_Config_Mandate'!$A:$A, 0))), TRUE, FALSE)"
    ws_runs[f"O{r}"].value = f"=IF($I{r} <= INDEX('02_Config_Mandate'!$E:$E, MATCH($E{r}, '02_Config_Mandate'!$A:$A, 0)), TRUE, FALSE)"
    ws_runs[f"P{r}"].value = f"=IF($K{r} <= INDEX('02_Config_Mandate'!$F:$F, MATCH($E{r}, '02_Config_Mandate'!$A:$A, 0)), TRUE, FALSE)"
    ws_runs[f"Q{r}"].value = f"=IF($J{r} <= INDEX('02_Config_Mandate'!$G:$G, MATCH($E{r}, '02_Config_Mandate'!$A:$A, 0)), TRUE, FALSE)"
    ws_runs[f"R{r}"].value = f"=IF($L{r} <= INDEX('02_Config_Mandate'!$H:$H, MATCH($E{r}, '02_Config_Mandate'!$A:$A, 0)), TRUE, FALSE)"
    ws_runs[f"S{r}"].value = f"=IF(AND($M{r}, $N{r}, $O{r}, $P{r}, $Q{r}, $R{r}), \"APPROVED\", \"FAIL_CLOSED\")"
    ws_runs[f"T{r}"].value = f"=TEXTJOIN(\"; \", TRUE, IF($M{r}, \"\", \"ERR_UNAUTHORIZED_MINT\"), IF($N{r}, \"\", \"ERR_VOLUME_EXCEEDED\"), IF($O{r}, \"\", \"ERR_STALE_ORACLE\"), IF($P{r}, \"\", \"ERR_CONFIDENCE_TOO_WIDE\"), IF($Q{r}, \"\", \"ERR_PRICE_DIVERGENCE\"), IF($R{r}, \"\", \"ERR_EXCESSIVE_SLIPPAGE\"))"
    ws_runs[f"U{r}"].value = f"=IF($S{r}=\"FAIL_CLOSED\", INDEX('04_Agent_Intent'!$C:$C, MATCH($G{r}, '04_Agent_Intent'!$A:$A, 0)), 0)"

wb.save(OUT_XLSX)
print(f"Arquivo gerado: {OUT_XLSX}")
